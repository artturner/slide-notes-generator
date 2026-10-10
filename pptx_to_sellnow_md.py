#!/usr/bin/env python3
"""
PowerPoint to Sellnow Markdown Converter

Converts PowerPoint presentations to specially formatted markdown files (Sellnow format).
Each slide is converted to a "--- SLIDE X ---" section with structured fields:
- Type: The type of slide (Title Slide, Content Slide, etc.)
- Title: The slide title
- Body: The main content with bullets and formatting
- Speaker Notes: Notes for the presenter
- Accessibility: Alt text and accessibility information
"""

import argparse
import re
import sys
from pathlib import Path
from pptx import Presentation


def detect_slide_type(slide, slide_index):
    """Detect the type of slide based on layout and position."""
    if slide_index == 0:
        return "Title Slide"

    layout_name = slide.slide_layout.name.lower()

    # Check layout names for common patterns
    if 'title' in layout_name and 'only' in layout_name:
        return "Section/Unit Slide"
    elif 'title' in layout_name and ('content' in layout_name or 'body' in layout_name):
        return "Content Slide"
    elif 'section' in layout_name:
        return "Section/Unit Slide"
    elif 'blank' in layout_name:
        return "Blank Slide"

    # Check if slide has content placeholder
    has_title = bool(slide.shapes.title)
    has_content = False

    for shape in slide.shapes:
        if hasattr(shape, "text_frame") and shape != slide.shapes.title:
            if shape.text_frame.text.strip():
                has_content = True
                break

    if has_title and has_content:
        return "Content Slide"
    elif has_title:
        return "Section/Unit Slide"
    else:
        return "Content Slide"


def extract_text_with_bullets(text_frame):
    """Extract text from a text frame, preserving bullet structure."""
    if not text_frame:
        return ""

    lines = []
    for paragraph in text_frame.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        level = paragraph.level

        # Format based on level (0 = main bullet, 1+ = sub-bullet)
        if level == 0:
            lines.append(f"* {text}")
        else:
            # Add 4 spaces per indentation level
            indent = "    " * level
            lines.append(f"{indent}* {text}")

    return "\n".join(lines)


def extract_body_content(slide, slide_type):
    """Extract body content from a slide."""
    body_parts = []

    # For title slides, look for subtitle
    if slide_type == "Title Slide":
        for shape in slide.shapes:
            if shape.is_placeholder:
                placeholder_format = shape.placeholder_format
                if placeholder_format.idx == 1:  # Subtitle placeholder
                    if hasattr(shape, "text_frame") and shape.text_frame.text.strip():
                        body_parts.append(shape.text_frame.text.strip())

    # For other slides, extract all text except title
    else:
        for shape in slide.shapes:
            # Skip title
            if hasattr(slide.shapes, 'title') and shape == slide.shapes.title:
                continue

            # Extract text from text frames
            if hasattr(shape, "text_frame"):
                content = extract_text_with_bullets(shape.text_frame)
                if content:
                    body_parts.append(content)

            # Extract text from tables
            elif hasattr(shape, "table"):
                table_text = []
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    table_text.append(row_text)
                if table_text:
                    body_parts.append("\n".join(table_text))

    return "\n\n".join(body_parts)


def extract_speaker_notes(slide):
    """Extract speaker notes from a slide."""
    if not slide.has_notes_slide:
        return ""

    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame

    if not text_frame:
        return ""

    # Combine all paragraphs into speaker notes
    notes = []
    for paragraph in text_frame.paragraphs:
        text = paragraph.text.strip()
        if text:
            notes.append(text)

    return " ".join(notes)


def generate_accessibility_text(slide, slide_type, title):
    """Generate basic accessibility information for a slide."""
    accessibility_lines = []

    # Count images
    image_count = sum(1 for shape in slide.shapes if shape.shape_type == 13)  # 13 = PICTURE

    if slide_type == "Title Slide":
        if image_count > 0:
            accessibility_lines.append(f"- **Alt Text for [Textbook Cover Image]:** {title} presentation cover.")
        else:
            accessibility_lines.append(f"- **Alt Text for [Slide Content]:** Title slide displaying: {title}.")

    elif image_count > 0:
        accessibility_lines.append(f"- **Alt Text for [Slide Content]:** Slide titled '{title}' with {image_count} image(s) and text content.")

    else:
        # Generic description
        if slide_type == "Section/Unit Slide":
            accessibility_lines.append(f"- **Alt Text for [Slide Content]:** Section header text: {title}.")
        else:
            accessibility_lines.append(f"- **Alt Text for [Slide Content]:** Slide content for '{title}'.")

    return "\n".join(accessibility_lines) if accessibility_lines else "- **Alt Text for [Slide Content]:** Slide content."


def convert_powerpoint_to_markdown(pptx_file):
    """Convert PowerPoint presentation to Sellnow markdown format."""
    prs = Presentation(pptx_file)
    markdown_lines = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        # Detect slide type
        slide_type = detect_slide_type(slide, slide_index - 1)

        # Extract title
        title = ""
        if slide.shapes.title:
            title = slide.shapes.title.text.strip()

        # Extract body
        body = extract_body_content(slide, slide_type)

        # Extract speaker notes
        speaker_notes = extract_speaker_notes(slide)

        # Generate accessibility text
        accessibility = generate_accessibility_text(slide, slide_type, title)

        # Format slide in Sellnow markdown
        markdown_lines.append(f"--- SLIDE {slide_index} ---")
        markdown_lines.append(f"**Type:** {slide_type}")
        markdown_lines.append(f"**Title:** {title}")
        markdown_lines.append("")
        markdown_lines.append("**Body:**")
        if body:
            markdown_lines.append(body)
        else:
            markdown_lines.append("[No content]")
        markdown_lines.append("")
        markdown_lines.append("**Speaker Notes:**")
        if speaker_notes:
            markdown_lines.append(speaker_notes)
        else:
            markdown_lines.append("[No speaker notes]")
        markdown_lines.append("")
        markdown_lines.append("**Accessibility:**")
        markdown_lines.append(accessibility)
        markdown_lines.append("")

    return "\n".join(markdown_lines)


def main():
    """Main function to handle command line arguments and execute conversion."""
    parser = argparse.ArgumentParser(
        description='Convert PowerPoint presentation to Sellnow-formatted markdown'
    )
    parser.add_argument(
        'input_file',
        help='Input PowerPoint file (.pptx)'
    )
    parser.add_argument(
        '-o', '--output',
        help='Output markdown file name (default: input_file.md)',
        default=None
    )

    args = parser.parse_args()

    # Check if input file exists
    input_path = Path(args.input_file)
    if not input_path.exists():
        print(f"Error: Input file '{args.input_file}' not found.", file=sys.stderr)
        sys.exit(1)

    # Check file extension
    if input_path.suffix.lower() != '.pptx':
        print(f"Error: Input file must be a .pptx file.", file=sys.stderr)
        sys.exit(1)

    # Determine output filename
    if args.output:
        output_file = args.output
    else:
        output_file = input_path.stem + '.md'

    try:
        # Convert PowerPoint to markdown
        print(f"Converting PowerPoint file: {args.input_file}")
        markdown_content = convert_powerpoint_to_markdown(input_path)

        # Write to output file
        output_path = Path(output_file)
        output_path.write_text(markdown_content, encoding='utf-8')

        # Count slides
        slide_count = markdown_content.count('--- SLIDE')
        print(f"Successfully converted {slide_count} slides")
        print(f"Markdown file saved as: {output_file}")
        print("Conversion completed successfully!")

    except Exception as e:
        print(f"Error: {str(e)}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
